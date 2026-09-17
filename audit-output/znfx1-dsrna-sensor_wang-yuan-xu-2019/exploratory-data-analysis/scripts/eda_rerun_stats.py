import os, glob, math, json
from collections import Counter
import openpyxl

# Source-data directory. Override with EDA_DATA_DIR; defaults to the copies
# archived in this repository under ../../supporting-files/source-data.
D = os.environ.get("EDA_DATA_DIR") or os.path.normpath(
    os.path.join(os.path.dirname(os.path.abspath(__file__)),
                 "..", "..", "supporting-files", "source-data"))

allvals = []
per_file = {}
for f in sorted(glob.glob(os.path.join(D, "*.xlsx"))):
    wb = openpyxl.load_workbook(f, data_only=True)
    vals, texts, blanks = [], 0, 0
    formulas = 0
    wb2 = openpyxl.load_workbook(f, data_only=False)
    for ws in wb2.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("="):
                    formulas += 1
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                v = c.value
                if v is None:
                    blanks += 1
                elif isinstance(v, (int, float)) and not isinstance(v, bool):
                    vals.append(float(v))
                else:
                    texts += 1
    per_file[os.path.basename(f)] = dict(n_numeric=len(vals), n_text=texts, n_blank=blanks,
                                         n_formula=formulas, vals=vals)
    allvals += vals

def decimals(x):
    s = repr(x)
    if "e" in s or "E" in s:
        return None
    return len(s.split(".")[1].rstrip("0")) if "." in s else 0

print("### Per-file cell inventory")
for k, v in per_file.items():
    vals = v["vals"]
    dd = Counter(decimals(x) for x in vals)
    ones = sum(1 for x in vals if x == 1.0)
    print(f"{k}: numeric={v['n_numeric']} text={v['n_text']} blank={v['n_blank']} formulas={v['n_formula']}")
    print(f"   min={min(vals):.4g} max={max(vals):.4g} exact_1.0={ones} ({100*ones/len(vals):.1f}%)")
    print(f"   decimal-place histogram: {dict(sorted(dd.items(), key=lambda t:(t[0] is None, t[0])))}")

print()
print("### Pooled terminal-digit analysis (values with >=2 decimals, non-integer)")
last = Counter()
penult = Counter()
first = Counter()
n = 0
for x in allvals:
    d = decimals(x)
    if d is None or d < 2 or x == 1.0:
        continue
    s = repr(x).split(".")[1].rstrip("0")
    last[s[-1]] += 1
    penult[s[-2]] += 1
    m = abs(x)
    while m < 1:
        m *= 10
    while m >= 10:
        m /= 10
    first[str(int(m))] += 1
    n += 1

def chisq(counter, keys, label):
    tot = sum(counter.get(k, 0) for k in keys)
    exp = tot / len(keys)
    x2 = sum((counter.get(k, 0) - exp) ** 2 / exp for k in keys)
    print(f"{label}: N={tot}  chi2({len(keys)-1})={x2:.2f}")
    print("   " + "  ".join(f"{k}:{counter.get(k,0)}" for k in keys))

digits = [str(i) for i in range(10)]
chisq(last, digits, "terminal digit  ")
chisq(penult, digits, "penultimate digit")

# Benford first-digit
tot = sum(first.values())
print(f"\nfirst significant digit (Benford), N={tot}")
bx2 = 0.0
for i in range(1, 10):
    obs = first.get(str(i), 0)
    exp = tot * math.log10(1 + 1 / i)
    bx2 += (obs - exp) ** 2 / exp
    print(f"   {i}: obs={obs:4d} ({100*obs/tot:5.1f}%)  benford={100*math.log10(1+1/i):5.1f}%")
print(f"   chi2(8)={bx2:.2f}")
