import openpyxl, json, sys
from pathlib import Path
base = Path(sys.argv[1])
out = {}
for f in sorted(base.glob("*.xlsx")):
    wb = openpyxl.load_workbook(f, data_only=True)
    sheets = {}
    for ws in wb.worksheets:
        rows = []
        for r in ws.iter_rows():
            vals = [(c.coordinate, c.value) for c in r if c.value is not None]
            if vals:
                rows.append(vals)
        sheets[ws.title] = {"dims": ws.dimensions, "max_row": ws.max_row, "max_col": ws.max_column, "rows": rows}
    out[f.name] = sheets
print(json.dumps(out, indent=1, default=str))
