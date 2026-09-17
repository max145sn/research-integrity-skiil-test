import os, glob, hashlib, json
import openpyxl

# Source-data directory. Override with EDA_DATA_DIR; defaults to the copies
# archived in this repository under ../../supporting-files/source-data.
D = os.environ.get("EDA_DATA_DIR") or os.path.normpath(
    os.path.join(os.path.dirname(os.path.abspath(__file__)),
                 "..", "..", "supporting-files", "source-data"))

for f in sorted(glob.glob(os.path.join(D, "*.xlsx"))):
    print("=" * 90)
    print("FILE:", os.path.basename(f), os.path.getsize(f), "bytes")
    wb = openpyxl.load_workbook(f, data_only=True)
    print("props:", wb.properties.creator, "|", wb.properties.created, "|", wb.properties.lastModifiedBy, "|", wb.properties.modified)
    print("sheets:", wb.sheetnames)
    for ws in wb.worksheets:
        print("-" * 80)
        print(f"SHEET '{ws.title}' dims={ws.dimensions} max_row={ws.max_row} max_col={ws.max_column}")
        rows = list(ws.iter_rows(values_only=True))
        for i, r in enumerate(rows):
            cells = ["" if v is None else (str(round(v, 6)) if isinstance(v, float) else str(v)) for v in r]
            while cells and cells[-1] == "":
                cells.pop()
            print(f"  r{i+1}: " + " | ".join(cells))
