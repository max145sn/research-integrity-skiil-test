#!/usr/bin/env python3
# /// script
# requires-python = ">=3.9"
# dependencies = ["openpyxl"]
# ///
"""Dump source-data sheets in a form you can actually compare against prose.

    uv run dump_sheet.py <workbook.xlsx> <sheet> [<sheet> ...]
    uv run dump_sheet.py <workbook.xlsx> --all-small     # every sheet under the row cap

Small sheets (comparison tables, per-reference bar charts) print in full, because
the whole point is to read every value. Large sheets (raw time-series traces)
print their header rows plus per-column statistics, because what you need from a
20,000-point trace is its mean and its settled plateau -- those are the numbers
that appear in the manuscript prose as "the output current increased to X mA" or
"the voltage stabilised at Y V".

The plateau statistic (last third of each column) matters more than it looks: a
trace that rises and then settles has a whole-column mean that matches nothing in
the paper, while its plateau matches the reported value exactly.
"""

import statistics
import sys
import pathlib

FULL_DUMP_ROW_CAP = 40
MAX_COLS = 24


def fmt(v):
    if v is None:
        return ""
    if isinstance(v, float):
        return round(v, 4)
    return v


def dump_sheet(ws, name: str) -> None:
    print("=" * 70)
    print(f"SHEET: {name}   dims={ws.calculate_dimension()}")
    print("=" * 70)

    rows = list(ws.iter_rows(max_col=MAX_COLS, values_only=True))
    nonempty = [r for r in rows if any(c is not None for c in r)]

    if len(nonempty) <= FULL_DUMP_ROW_CAP:
        for i, row in enumerate(rows, 1):
            vals = [fmt(v) for v in row]
            while vals and str(vals[-1]) == "":
                vals.pop()
            if vals:
                print(f"{i:>4} {vals}")
        print()
        return

    print(f"-- header rows (first 4 of {len(nonempty)} non-empty) --")
    shown = 0
    for i, row in enumerate(rows, 1):
        vals = [fmt(v) for v in row]
        if any(str(v) != "" for v in vals):
            print(f"{i:>4} {vals}")
            shown += 1
            if shown >= 4:
                break

    cols = {}
    for row in rows:
        for j, v in enumerate(row):
            if isinstance(v, (int, float)) and not isinstance(v, bool):
                cols.setdefault(j, []).append(v)

    print(f"-- numeric column stats (col: n, min, mean, max | plateau=mean of last third) --")
    for j, vals in sorted(cols.items()):
        if len(vals) < 3:
            continue
        tail = vals[len(vals) * 2 // 3:] or vals
        print(
            f"  col{j:<3} n={len(vals):<7} min={min(vals):<12.4g} "
            f"mean={statistics.mean(vals):<12.4g} max={max(vals):<12.4g} "
            f"plateau={statistics.mean(tail):.4g}"
        )
    print()


def main() -> int:
    if len(sys.argv) < 3:
        print(__doc__)
        return 2

    import openpyxl

    path = pathlib.Path(sys.argv[1]).expanduser().resolve()
    wanted = sys.argv[2:]
    wb = openpyxl.load_workbook(path, read_only=True)

    if wanted == ["--all-small"]:
        targets = wb.sheetnames
    else:
        targets = []
        for w in wanted:
            match = next((s for s in wb.sheetnames if s.lower() == w.lower()), None)
            if match is None:
                print(f"!! no sheet named {w!r}. Available: {wb.sheetnames}", file=sys.stderr)
                continue
            targets.append(match)

    for name in targets:
        ws = wb[name]
        if wanted == ["--all-small"]:
            rows = list(ws.iter_rows(max_col=MAX_COLS, values_only=True))
            if len([r for r in rows if any(c is not None for c in r)]) > FULL_DUMP_ROW_CAP:
                continue
        dump_sheet(ws, name)

    wb.close()
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
