#!/usr/bin/env python3
# /// script
# requires-python = ">=3.9"
# dependencies = ["openpyxl"]
# ///
"""Ask of a source-data workbook: are these numbers independent measurements?

    uv run scan_workbook.py <workbook.xlsx>
    uv run scan_workbook.py <workbook.xlsx> --digits   # add the terminal-digit tally

Four passes. The first carries findings on its own; the rest are pointers that
have to resolve into something demonstrable before they are worth reporting.

1. DERIVED REPLICATES. openpyxl's data_only=True returns cached results and
   discards every formula, so a cell holding "=3-C70-D70" reads back as a
   measured 0.939 and nothing looks wrong. This pass loads the book both ways.
   A formula in a mean or an SD is expected. A formula in a cell the figure
   plots as one of n independent replicates is not: that value is fixed by its
   neighbours, so the group has n-1 free values, its mean is pinned to a
   constant, and the SD summarises n-1 numbers. Where a legend claims "n = 3
   independent experiments", the workbook is then supplying two.

2. DUPLICATE BLOCKS. Identical runs of values in two places, at full precision.

3. INTEGER STEPS. Groups whose values differ from the row above by whole numbers
   while parallel groups on the same row differ by ordinary amounts.

4. TERMINAL DIGITS (--digits). Last-digit concentration across a sheet.

Report the artifact and the arithmetic, never a motive. The finding is that the
effective replicate count is below the stated one; the ask is for the raw
measurements.
"""

import pathlib
import re
import sys
from collections import Counter, defaultdict

CONST_MINUS = re.compile(r"^=\s*(\d+(?:\.\d+)?)\s*-\s*[A-Z$]{1,3}\$?\d+(\s*-\s*[A-Z$]{1,3}\$?\d+)*\s*$")
SUMMARY_FN = re.compile(r"\b(AVERAGE|STDEV|MEDIAN|SUM|COUNT|VAR|SEM)\b", re.I)
MIN_INTEGER_RUN = 4


def col_letter(idx):
    s = ""
    while idx > 0:
        idx, r = divmod(idx - 1, 26)
        s = chr(65 + r) + s
    return s


def numeric_runs(ws, row, max_col):
    """Contiguous horizontal runs of numbers -- the shape a replicate group takes."""
    runs, cur, start = [], [], None
    for c in range(1, max_col + 2):
        v = ws.cell(row, c).value if c <= max_col else None
        if isinstance(v, (int, float)) and not isinstance(v, bool):
            if start is None:
                start = c
            cur.append((c, v))
        else:
            if len(cur) >= 2:
                runs.append((start, cur))
            cur, start = [], None
    return runs


def scan_formulas(wbf, wbv):
    print("=" * 72)
    print("1. DERIVED REPLICATES  (what data_only=True would have hidden)")
    print("=" * 72)
    shapes, cells = Counter(), []
    for wsf in wbf.worksheets:
        wsv = wbv[wsf.title]
        for row in wsf.iter_rows():
            for cell in row:
                f = cell.value
                if not (isinstance(f, str) and f.startswith("=")):
                    continue
                shapes[re.sub(r"[A-Z$]{1,3}\$?\d+", "REF", f)] += 1
                cells.append((wsf.title, wsv, cell, f))

    if not cells:
        print("  none -- every cell is a static value.\n")
        return
    print(f"  {len(cells)} formula cell(s) in {len({c[0] for c in cells})} sheet(s)\n")
    print("  shapes:")
    for shape, n in shapes.most_common():
        print(f"    {n:5d}  {shape}")

    print("\n  classified:")
    for sheet, wsv, cell, f in cells:
        val = wsv.cell(cell.row, cell.column).value
        print(f"    {sheet}!{cell.coordinate} = {f}   -> {val!r}")
        if SUMMARY_FN.search(f):
            print("        expected -- summary statistic, not a data point")
            continue
        refs = re.findall(r"[A-Z$]{1,3}\$?\d+", f)
        deps = []
        for ref in refs:
            rc = re.match(r"\$?([A-Z]{1,3})\$?(\d+)", ref)
            deps.append(wsv[f"{rc.group(1)}{rc.group(2)}"].value)
        m = CONST_MINUS.match(f.replace(" ", ""))
        if m and all(isinstance(d, (int, float)) for d in deps):
            const = float(m.group(1))
            group = [val] + deps
            print(f"        REPORTABLE -- this replicate is defined as {const:g} minus the "
                  f"other {len(deps)}.")
            print(f"        group {[round(x, 6) for x in group]} sums to {sum(group):g}; "
                  f"mean pinned to {const / len(group):g}")
            print(f"        effective n = {len(deps)}, not {len(group)}: the group has "
                  f"{len(deps)} free values, so the SD summarises {len(deps)} numbers.")
        else:
            print(f"        check -- formula depends on {len(refs)} other cell(s); confirm "
                  f"whether the figure plots this as an independent replicate")
    print()


def scan_duplicates(wbv, width=3):
    print("=" * 72)
    print(f"2. DUPLICATED BLOCKS  (identical runs of {width}, full precision)")
    print("=" * 72)
    seen = defaultdict(list)
    for ws in wbv.worksheets:
        for r in range(1, ws.max_row + 1):
            for start, run in numeric_runs(ws, r, ws.max_column):
                vals = [v for _, v in run]
                for k in range(len(vals) - width + 1):
                    key = tuple(vals[k:k + width])
                    if len(set(key)) > 1:
                        seen[key].append(f"{ws.title}!{col_letter(run[k][0])}{r}")
    dups = {k: sorted(set(v)) for k, v in seen.items() if len(set(v)) > 1}
    if not dups:
        print("  none.\n")
        return
    for key, locs in sorted(dups.items(), key=lambda kv: -max(len(repr(x)) for x in kv[0])):
        print(f"  {[round(x, 10) for x in key]}")
        for loc in locs:
            print(f"      {loc}")
    print()


def scan_integer_steps(wbv):
    print("=" * 72)
    print("3. INTEGER STEPS BETWEEN ROWS  (a group stepping by whole numbers)")
    print("=" * 72)
    found = False
    for ws in wbv.worksheets:
        prev = None
        for r in range(1, ws.max_row + 1):
            runs = numeric_runs(ws, r, ws.max_column)
            cur = {start: [v for _, v in run] for start, run in runs}
            if prev:
                for start, vals in cur.items():
                    old = prev.get(start)
                    if not old or len(old) != len(vals):
                        continue
                    # a whole row is often one long run holding several groups
                    # side by side, so look for integer-stepping WINDOWS inside it
                    ok = [abs((b - a) - round(b - a)) < 1e-9 and abs(b - round(b)) > 1e-9
                          for a, b in zip(old, vals)]
                    i = 0
                    while i < len(ok):
                        if not ok[i]:
                            i += 1
                            continue
                        j = i
                        while j < len(ok) and ok[j]:
                            j += 1
                        if j - i >= MIN_INTEGER_RUN:
                            found = True
                            print(f"  {ws.title} row {r}, cols "
                                  f"{col_letter(start + i)}-{col_letter(start + j - 1)} "
                                  f"vs the row above")
                            print(f"      from  {[round(v, 6) for v in old[i:j]]}")
                            print(f"      to    {[round(v, 6) for v in vals[i:j]]}")
                            print(f"      steps {[round(b - a, 6) for a, b in zip(old[i:j], vals[i:j])]}")
                        i = j
            prev = cur
    if not found:
        print("  none.\n")
    else:
        print()


def scan_digits(wbv):
    print("=" * 72)
    print("4. TERMINAL DIGITS  (pointer only -- never a finding on its own)")
    print("=" * 72)
    for ws in wbv.worksheets:
        tally = Counter()
        for row in ws.iter_rows(values_only=True):
            for v in row:
                if isinstance(v, float) and not isinstance(v, bool):
                    s = repr(round(v, 10)).rstrip("0")
                    if "." in s and s[-1].isdigit():
                        tally[s[-1]] += 1
        if sum(tally.values()) < 50:
            continue
        n = sum(tally.values())
        bars = "  ".join(f"{d}:{100 * tally[d] / n:4.1f}%" for d in "0123456789")
        print(f"  {ws.title:26s} n={n:<6d} {bars}")
    print()


def main():
    if len(sys.argv) < 2:
        print(__doc__)
        return 2
    import openpyxl

    path = pathlib.Path(sys.argv[1]).expanduser().resolve()
    wbf = openpyxl.load_workbook(path, data_only=False)
    wbv = openpyxl.load_workbook(path, data_only=True)
    print(f"workbook: {path.name}\nsheets:   {len(wbv.worksheets)}\n")
    scan_formulas(wbf, wbv)
    scan_duplicates(wbv)
    scan_integer_steps(wbv)
    if "--digits" in sys.argv:
        scan_digits(wbv)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
